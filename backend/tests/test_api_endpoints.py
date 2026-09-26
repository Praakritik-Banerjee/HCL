import io
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
import docx

from app.main import app
from app.db.base import Base
from app.db.session import get_db
from sqlalchemy.pool import StaticPool

# Setup test in-memory database with StaticPool so all connections share the same tables
test_engine = create_engine(
    "sqlite:///:memory:",
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)
Base.metadata.create_all(bind=test_engine)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=test_engine)


def override_get_db():
    db = TestingSessionLocal()
    try:
        yield db
    finally:
        db.close()


app.dependency_overrides[get_db] = override_get_db
client = TestClient(app)


def test_root_endpoint():
    response = client.get("/")
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "online"
    assert "docs" in data


def test_error_shape_on_invalid_upload():
    # Attempt upload with invalid file extension (.exe instead of .pdf/.docx/.txt)
    files = {"file": ("notes.exe", b"executable data", "application/octet-stream")}
    response = client.post("/api/v1/syllabus/upload", files=files)


    assert response.status_code == 400
    data = response.json()
    # Verify exact required error shape
    assert "error_code" in data
    assert "message" in data
    assert data["error_code"] == "UNSUPPORTED_FILE_TYPE"


def test_syllabus_upload_and_graph_query():
    # Build in-memory DOCX
    doc = docx.Document()
    doc.add_heading("Course Syllabus: Data Structures", level=0)
    doc.add_heading("Unit 1: Linear Structures", level=1)
    doc.add_paragraph("Arrays, dynamic arrays, linked lists, stacks and queues.")
    doc.add_heading("1.1 Dynamic Arrays", level=2)
    doc.add_paragraph("Amortized analysis of array doubling, pointer manipulation.")
    doc.add_heading("Unit 2: Non-Linear Structures", level=1)
    doc.add_paragraph("Binary search trees, AVL trees, graphs and traversals.")

    docx_io = io.BytesIO()
    doc.save(docx_io)
    docx_bytes = docx_io.getvalue()

    # 1. Upload
    upload_resp = client.post(
        "/api/v1/syllabus/upload",
        files={"file": ("data_structures_syllabus.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )

    assert upload_resp.status_code == 201
    upload_data = upload_resp.json()
    assert upload_data["success"] is True
    doc_id = upload_data["data"]["document_id"]
    assert upload_data["data"]["total_topics"] >= 2
    assert upload_data["data"]["status"] == "completed"

    # 2. Query Knowledge Graph
    graph_resp = client.get(f"/api/v1/syllabus/{doc_id}/graph")
    assert graph_resp.status_code == 200
    graph_data = graph_resp.json()
    assert graph_data["success"] is True
    assert graph_data["data"]["document_id"] == doc_id
    assert graph_data["data"]["total_nodes"] >= 2
    assert len(graph_data["data"]["topics"]) >= 2

    # Verify hierarchical structure
    unit_nodes = graph_data["data"]["topics"]
    assert any("Linear Structures" in node["title"] for node in unit_nodes)

    # 3. Query 404 on non-existent document
    not_found_resp = client.get("/api/v1/syllabus/non-existent-uuid/graph")
    assert not_found_resp.status_code == 404
    err = not_found_resp.json()
    assert err["error_code"] == "RESOURCE_NOT_FOUND"
