import io
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
import docx

from app.db.base import Base
from app.models.document import Document
from app.models.knowledge_graph import Topic, TopicChunk
from app.services.ingestion.cleaner import TextCleaner
from app.services.ingestion.graph_builder import KnowledgeGraphBuilder
from app.services.ingestion.chunker import HierarchicalChunker, count_tokens
from app.services.ingestion.pipeline import IngestionPipeline
from tests.create_sample_syllabus import SAMPLE_SYLLABUS_TEXT


@pytest.fixture
def db_session():
    """Create an isolated in-memory SQLite database session for unit testing."""
    test_engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False})
    Base.metadata.create_all(bind=test_engine)
    TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=test_engine)
    session = TestingSessionLocal()
    yield session
    session.close()


def test_pii_scrubbing_and_cleaning():
    raw_sample = (
        "Contact Prof. Alan Turing at alan.turing@cambridge.ac.uk or +1-555-019-2831. "
        "Student ID: STU-1029384.\n"
        "Page 1 of 15\n"
        "Course Overview"
    )
    cleaned = TextCleaner.clean(raw_sample)

    # Assert PII is redacted
    assert "alan.turing@cambridge.ac.uk" not in cleaned
    assert "[REDACTED_EMAIL]" in cleaned
    assert "+1-555-019-2831" not in cleaned
    assert "[REDACTED_PHONE]" in cleaned
    assert "STU-1029384" not in cleaned
    assert "[REDACTED_ID]" in cleaned
    # Assert page number line removed
    assert "Page 1 of 15" not in cleaned


def test_knowledge_graph_hierarchy_extraction():
    parsed_units = KnowledgeGraphBuilder.parse_hierarchy(SAMPLE_SYLLABUS_TEXT)

    # Acceptance Criteria: Captures 100% of top-level units
    assert len(parsed_units) == 4
    assert parsed_units[0].title == "Foundations of Supervised Learning"
    assert parsed_units[0].unit_label == "Unit 1"
    assert parsed_units[1].title == "Deep Learning Architectures"
    assert parsed_units[2].title == "Unsupervised Learning & Dimensionality Reduction"
    assert parsed_units[3].title == "Reinforcement Learning & Autonomous Agents"

    # Verify nested subtopics under Unit 1
    u1_children = parsed_units[0].children
    assert len(u1_children) == 3
    assert "Linear Regression and Gradient Descent" in u1_children[0].title
    assert "Logistic Regression and Classification Metrics" in u1_children[1].title
    assert "Regularization and the Bias-Variance Tradeoff" in u1_children[2].title


def test_chunker_token_bounds_and_ancestry():
    chunker = HierarchicalChunker(max_tokens=400, min_tokens=100)
    section_text = (
        "Supervised learning algorithms build a mathematical model of a set of data that contains both the inputs and the desired outputs. "
        "The data is known as training data, and consists of a set of training examples. "
        "Each training example has one or more inputs and the desired output, also known as a supervisory signal. "
        "In the case of semi-supervised algorithms, some of the training examples are missing the desired output. "
        "Linear regression models target a continuous numerical output value, optimizing the residual sum of squares."
    )
    chunks = chunker.chunk_section(
        text=section_text,
        heading_path="Unit 1 > 1.1 Linear Regression",
        topic_title="Linear Regression",
        level=2,
    )

    assert len(chunks) >= 1
    for chunk in chunks:
        assert chunk.token_count <= 400
        assert chunk.heading_path == "Unit 1 > 1.1 Linear Regression"
        assert chunk.topic_title == "Linear Regression"
        assert chunk.level == 2
        assert len(chunk.content) > 0


def test_ingestion_pipeline_end_to_end_docx(db_session):
    # Synthesize a real .docx binary file in memory
    doc = docx.Document()
    doc.add_heading("CS 480 Machine Learning Syllabus", level=0)
    doc.add_heading("Unit 1: Foundations of Supervised Learning", level=1)
    doc.add_paragraph("Overview of basic learning algorithms and optimization methods.")
    doc.add_heading("1.1 Linear Regression", level=2)
    doc.add_paragraph("Empirical risk minimization, MSE loss, and stochastic gradient descent.")
    doc.add_heading("Unit 2: Neural Networks", level=1)
    doc.add_paragraph("Backpropagation and deep feedforward networks.")

    docx_bytes_io = io.BytesIO()
    doc.save(docx_bytes_io)
    docx_bytes = docx_bytes_io.getvalue()

    pipeline = IngestionPipeline(db=db_session)
    result = pipeline.process_file(
        file_bytes=docx_bytes,
        filename="test_syllabus.docx",
        file_type="docx",
    )

    assert result.status == "completed"
    assert result.total_topics >= 2
    assert result.total_chunks >= 2

    # Query PostgreSQL/SQLite to verify relational knowledge graph
    stored_doc = db_session.query(Document).filter(Document.id == result.document_id).first()
    assert stored_doc is not None
    assert stored_doc.filename == "test_syllabus.docx"

    topics = db_session.query(Topic).filter(Topic.document_id == result.document_id).all()
    assert len(topics) >= 2

    unit_titles = [t.title for t in topics if t.level == 1]
    assert any("Supervised Learning" in t for t in unit_titles)
    assert any("Neural Networks" in t for t in unit_titles)

    chunks = db_session.query(TopicChunk).filter(TopicChunk.document_id == result.document_id).all()
    assert len(chunks) >= 2
    for chk in chunks:
        assert chk.vector_id.startswith("chk_")
        assert chk.token_count > 0


def test_ingestion_pipeline_end_to_end_pdf(db_session):
    from tests.create_sample_syllabus import create_sample_pdf_bytes

    pdf_bytes = create_sample_pdf_bytes()
    pipeline = IngestionPipeline(db=db_session)
    result = pipeline.process_file(
        file_bytes=pdf_bytes,
        filename="ai_ml_syllabus.pdf",
        file_type="pdf",
    )

    assert result.status == "completed"
    assert result.total_topics == 4
    assert result.total_subtopics >= 8
    assert result.total_chunks >= 4

    stored_doc = db_session.query(Document).filter(Document.id == result.document_id).first()
    assert stored_doc is not None
    assert stored_doc.filename == "ai_ml_syllabus.pdf"

    # Verify all 4 units were persisted
    units = db_session.query(Topic).filter(Topic.document_id == result.document_id, Topic.level == 1).all()
    assert len(units) == 4
    titles = [u.title for u in units]
    assert any("Foundations of Supervised Learning" in t for t in titles)
    assert any("Deep Learning Architectures" in t for t in titles)
    assert any("Unsupervised Learning" in t for t in titles)
    assert any("Reinforcement Learning" in t for t in titles)

